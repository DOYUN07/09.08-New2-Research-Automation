"""파서 검증 — 네트워크 없이 저장된 마크업으로 확인한다.

실행: python -m tests.test_parse
"""
from __future__ import annotations

import sys
from datetime import date

from src.config import Institution, load_config, load_institutions
from src.filters import FilterStats, apply_filters
from src.parse import parse
from src.verify import drop_expired as drop_expired_fn
from tests import fixtures

TODAY = date(2026, 9, 7)
FAILS: list[str] = []


def check(label: str, got, want) -> None:
    if got != want:
        FAILS.append(f"{label}\n    기대: {want!r}\n    실제: {got!r}")
        print(f"  ✗ {label}: {got!r} != {want!r}")
    else:
        print(f"  ✓ {label}")


def inst(**kw) -> Institution:
    kw.setdefault("id", "t")
    kw.setdefault("name", "테스트")
    kw.setdefault("url", "https://example.com/list")
    kw.setdefault("base", "https://example.com")
    return Institution(**kw)


def load_institutions_for_test():
    return load_institutions()


def P(html: str, institution: Institution):
    """기준일을 2026-09-07로 고정해서 파싱한다 (미래 날짜 판정이 결과에 영향을 주므로)."""
    return parse(html, institution, TODAY)


def main() -> int:
    print("\n[1] 부산테크노파크형 — 접수기간과 게시일이 따로 있는 표")
    r = P(fixtures.BTP, inst(id="btp", base="https://www.btp.or.kr/kor/CMS/Board/Board.do"))
    check("행 수", len(r), 4)
    check("제목", r[0].title, "Age-Tech 종합지원센터 운영 사업 TRL기반 기술성장 맞춤 지원 공고")
    check("게시일", r[0].posted, date(2026, 9, 7))
    check("마감일", r[0].deadline, date(2026, 9, 30))
    check("링크 절대경로", r[0].url.startswith("https://www.btp.or.kr/"), True)
    check("마감 상태 인식", r[1].closed_flag, True)

    print("\n[2] 광주테크노파크형 — 기간만 있고 게시일 컬럼 없음, 번호가 th")
    r = P(fixtures.GJTP, inst(id="gjtp", base="https://www.gjtp.or.kr/home/business.cs"))
    check("행 수", len(r), 3)
    check("게시일=기간 시작", r[0].posted, date(2026, 9, 3))
    check("마감일=기간 끝", r[0].deadline, date(2026, 9, 14))

    print("\n[3] 농림축산식품부형 — 날짜가 dd.date 안")
    r = P(fixtures.MAFRA, inst(id="mafra", base="https://www.mafra.go.kr"))
    check("행 수", len(r), 3)
    check("제목에서 '새글' 제거", r[0].title, "스마트농업 클라우드 실증 지원사업 공고")
    check("게시일", r[0].posted, date(2026, 9, 4))

    print("\n[4] IRIS형 — ul/li 구조 + onclick 링크")
    r = P(
        fixtures.IRIS,
        inst(
            id="iris",
            base="https://www.iris.go.kr",
            row_selector="ul.dbody > li",
            link_from_onclick=r"_view\('(\d+)'",
            detail_url="https://www.iris.go.kr/contents/retrieveBsnsAncmView.do?ancmId={id}&ancmPrg=ancmIng",
        ),
    )
    check("행 수", len(r), 3)
    check("게시일", r[0].posted, date(2026, 9, 7))
    check(
        "onclick에서 상세 URL 조립",
        r[0].url,
        "https://www.iris.go.kr/contents/retrieveBsnsAncmView.do?ancmId=023977&ancmPrg=ancmIng",
    )

    print("\n[4-b] IRIS형 — 선택자 없이 자동 인식되는지")
    r_auto = P(
        fixtures.IRIS,
        inst(
            id="iris",
            base="https://www.iris.go.kr",
            link_from_onclick=r"_view\('(\d+)'",
            detail_url="https://www.iris.go.kr/x?ancmId={id}",
        ),
    )
    check("자동 인식 행 수", len(r_auto), 3)

    print("\n[5] 중소벤처기업부형 — onclick doBbsFView + 중첩 div 안의 신청기간")
    r = P(
        fixtures.MSS,
        inst(
            id="mss",
            base="https://www.mss.go.kr",
            link_from_onclick=r"doBbsFView\('\d+','(\d+)'",
            detail_url="https://www.mss.go.kr/site/smba/ex/bbs/View.do?cbIdx=310&bcIdx={id}&parentSeq={id}",
        ),
    )
    check("행 수", len(r), 3)
    check("게시일", r[0].posted, date(2026, 9, 7))
    check("마감일=신청기간 종료", r[0].deadline, date(2026, 10, 6))
    check("상세 URL bcIdx", "bcIdx=1071012" in r[0].url, True)

    print("\n[6] 부산사회서비스원형 — 2자리 연도(26.09.05)")
    r = P(fixtures.BUSAN_PASS, inst(id="bp", base="https://busan.pass.or.kr/SW_bbs/notice/"))
    check("행 수", len(r), 3)
    check("2자리 연도 해석", r[0].posted, date(2026, 9, 5))

    print("\n[7] K-Startup형 — 카드형 목록 + go_view(id)")
    r = P(
        fixtures.KSTARTUP,
        inst(
            id="ks",
            base="https://www.k-startup.go.kr",
            link_from_onclick=r"go_view\((\d+)\)",
            detail_url="https://www.k-startup.go.kr/web/contents/bizpbanc-ongoing.do?schM=view&pbancSn={id}&page=1",
        ),
    )
    check("행 수", len(r), 3)
    check("게시일", r[0].posted, date(2026, 9, 4))
    check("마감일", r[0].deadline, date(2026, 9, 15))
    check("상세 URL", "pbancSn=179111" in r[0].url, True)

    print("\n[8] 게시일이 아예 없는 게시판")
    r = P(fixtures.NO_DATE, inst(id="nd", base="https://www.khidi.or.kr"))
    check("행 수", len(r), 3)
    check("게시일 None", r[0].posted, None)
    check("제목", r[0].title, "[공고] 2026년 1차 고령친화우수제품 지정 공고")

    # -------------------------------- 1차 실전 진단(2026-09-07)에서 발견된 문제들
    print("\n[A] 첨부파일 링크를 제목으로 착각하지 않는지 (산업통상부에서 발생)")
    r = P(fixtures.ATTACH_TRAP, inst(id="motir", base="https://www.motir.go.kr"))
    check("행 수", len(r), 3)
    check("제목이 파일명이 아님", r[0].title, "2026년도 산업통상부-에너지공기업 기술나눔")
    check("'.hwpx' 미포함", ".hwpx" in r[0].title, False)
    check("'다운로드' 미포함", "다운로드" in r[0].title, False)
    check("링크가 첨부파일이 아님", "download" in r[0].url, False)
    check("링크가 본문", "/71310/view" in r[0].url, True)

    print("\n[B] 제목이 두 번 이어붙는 현상 (부산테크노파크에서 발생)")
    r = P(fixtures.DOUBLED_TITLE, inst(id="dbl", base="https://example.com"))
    check("행 수", len(r), 3)
    check("제목 1회만", r[0].title, "신중년 디지털 전환 지원사업 공고")
    check("제목 1회만 (2)", r[1].title, "고령친화 서비스 실증 참여기업 모집")

    print("\n[C] 한 행에 기간이 두 개일 때 게시일이 미래로 잡히지 않는지 (연구개발특구진흥재단)")
    r = P(fixtures.TWO_PERIODS, inst(id="innopolis", base="https://pms.innopolis.or.kr"))
    check("행 수", len(r), 3)
    check("게시일=공고기간 시작", r[0].posted, date(2026, 8, 21))
    check("게시일이 오늘 이전", r[0].posted <= TODAY, True)
    check("마감일=가장 늦은 종료일", r[0].deadline, date(2026, 9, 21))
    check("2행 게시일", r[1].posted, date(2026, 9, 1))
    check("2행 마감일", r[1].deadline, date(2026, 10, 5))

    # -------------------------------- 2차 실전 진단(2026-09-07)에서 발견된 문제들
    print("\n[D] 목록에 마감일시만 있을 때 (전북테크노파크)")
    r = P(fixtures.DEADLINE_ONLY, inst(id="jbtp", base="https://jbcis.jbtp.or.kr"))
    check("행 수", len(r), 3)
    check("마감일시를 게시일로 쓰지 않음", r[0].posted, None)
    check("마감일로 인식", r[0].deadline, date(2026, 12, 31))
    check("3행 마감일", r[2].deadline, date(2026, 9, 30))

    print("\n[E] 두 span 내용이 완전히 같지 않은 중복 제목 (부산테크노파크)")
    r = P(fixtures.PARTIAL_DOUBLE, inst(id="btp2", base="https://www.btp.or.kr"))
    check("행 수", len(r), 3)
    check(
        "긴 쪽만 채택",
        r[0].title,
        "기업성장기반 글로벌 하이메디 허브 특구 상생협력사업 기업지원모집 공고(4차) 재공고",
    )
    check("말줄임표 미포함", "..." in r[0].title, False)
    check("2행도 1회만", r[1].title.count("글로벌시장"), 1)
    check("완전 동일한 경우도 1회", r[2].title, "시니어 돌봄로봇 실증 참여기업 모집 공고")

    print("\n[F] 게시판이 <header> 안에 있어도 인식 (NIPA 0건 회귀 방지)")
    r = P(fixtures.INSIDE_HEADER, inst(id="nipa", base="https://www.nipa.kr"))
    check("행 수", len(r), 3)
    check("제목", r[0].title, "2026년 아태 AI 특화지구(AHAP) 조성 사업 공고")
    check("class에 down이 있어도 링크 유지", "/home/2-2/16921" in r[0].url, True)
    check("게시일", r[0].posted, date(2026, 9, 3))

    # -------------------------------- 3차 실전 진단(2026-09-07)에서 발견된 문제들
    print("\n[G] 행 안에 첨부파일 <li>가 있어도 인식 (중소벤처기업부)")
    r = P(
        fixtures.ROW_WITH_LI,
        inst(
            id="mss",
            base="https://www.mss.go.kr",
            link_from_onclick=r"doBbsFView\('\d+','(\d+)'",
            detail_url="https://www.mss.go.kr/site/smba/ex/bbs/View.do?cbIdx=310&bcIdx={id}&parentSeq={id}",
        ),
    )
    check("행 수", len(r), 3)
    check(
        "제목에 메타데이터가 안 붙음",
        r[0].title,
        "2026년 중소기업 스마트서비스 지원사업 참여기업 모집 공고(A/S지원)",
    )
    check("'담당부서' 미포함", "담당부서" in r[0].title, False)
    check("게시일", r[0].posted, date(2026, 9, 7))
    check("마감일=신청기간 종료", r[0].deadline, date(2026, 10, 6))
    check("상세 URL", "bcIdx=1071012" in r[0].url, True)

    print("\n[H] 행마다 감싸는 div가 따로 있어 형제가 아닌 목록 (부산시민운동지원센터)")
    r = P(fixtures.WRAPPED_ROWS, inst(id="ng", base="https://www.ngocenter.or.kr/business/"))
    check("행 수", len(r), 3)
    check("제목", r[0].title, "[활동가커뮤니티지원사업] 든든 커뮤니티 큰모임 (9/12)")
    check("게시일", r[0].posted, date(2026, 9, 2))
    check("상대경로 결합", r[0].url, "https://www.ngocenter.or.kr/business/view?scti=0&no=2904")
    check("헤더 행 제외", any("제목" == n.title for n in r), False)

    # ---------------------------------------------------------------- 필터
    print("\n[9] 필터 규칙")
    cfg = load_config()
    cfg.include_keywords = ["Age-Tech", "지원사업", "고령", "AX", "실증", "돌봄", "클라우드"]
    cfg.exclude_keywords = ["채용", "포상", "구제역", "퇴직공무원"]
    cfg.lookback_days = 10
    cfg.max_per_institution = 10

    parsed = P(fixtures.BTP, inst(id="btp", base="https://www.btp.or.kr/"))
    st = FilterStats()
    kept = apply_filters(parsed, cfg, TODAY, set(), st)
    titles = [n.title for n in kept]
    check("채용 공고 제외", any("채용" in t for t in titles), False)
    check("마감된 공고 제외", any("커피산업" in t for t in titles), False)
    check("Age-Tech 공고 채택", any("Age-Tech" in t for t in titles), True)

    print("\n[9-b] 게시일 10일 초과 공고 제외")
    parsed = P(fixtures.BUSAN_PASS, inst(id="bp", base="https://busan.pass.or.kr/"))
    cfg.include_keywords = ["사회서비스", "돌봄"]
    cfg.exclude_keywords = []
    st = FilterStats()
    kept = apply_filters(parsed, cfg, TODAY, set(), st)
    check("07.29 공고 제외됨", any("역량강화" in n.title for n in kept), False)
    check("09.05 / 09.02 공고 채택", len(kept), 2)

    print("\n[9-c] 중복 발송 방지")
    st = FilterStats()
    already = {kept[0].key}
    kept2 = apply_filters(
        P(fixtures.BUSAN_PASS, inst(id="bp", base="https://busan.pass.or.kr/")),
        cfg,
        TODAY,
        already,
        st,
    )
    check("이미 보낸 공고 제외", len(kept2), 1)
    check("중복 카운트", st.dropped_duplicate, 1)

    print("\n[9-d] 마감일 미상 공고는 '확인필요'로 포함")
    cfg.include_keywords = ["고령", "에이지테크", "지정"]
    st = FilterStats()
    kept3 = apply_filters(
        P(fixtures.NO_DATE, inst(id="nd", base="https://www.khidi.or.kr")),
        cfg,
        TODAY,
        set(),
        st,
    )
    check("마감일 미상도 채택됨", len(kept3) >= 2, True)
    check("마감일 None", kept3[0].deadline, None)

    print("\n[9-e] unknown_deadline: exclude 로 바꾸면 제외")
    cfg.unknown_deadline = "exclude"
    st = FilterStats()
    kept4 = apply_filters(
        P(fixtures.NO_DATE, inst(id="nd", base="https://www.khidi.or.kr")),
        cfg,
        TODAY,
        set(),
        st,
    )
    check("전부 제외", len(kept4), 0)

    # ---------------------------------------------------------------- 누적 기록
    print("\n[10] 누적 기록 (엑셀)")
    import tempfile
    from pathlib import Path

    from src import archive

    src_notices = P(fixtures.BTP, inst(id="btp", name="부산테크노파크", base="https://www.btp.or.kr/"))
    for n in src_notices:
        n.matched_keywords = ["테스트"]

    with tempfile.TemporaryDirectory() as td:
        csvp = Path(td) / "archive.csv"
        xlsxp = Path(td) / "공고누적.xlsx"

        rows, added = archive.add([], src_notices, date(2026, 9, 7))
        check("1일차 추가 건수", added, len(src_notices))
        archive.save_csv(rows, csvp)
        check("CSV 저장됨", csvp.exists(), True)

        # 같은 공고를 다시 넣어도 늘지 않아야 한다
        rows2 = archive.load_rows(csvp)
        check("CSV 다시 읽기", len(rows2), len(src_notices))
        rows2, added2 = archive.add(rows2, src_notices, date(2026, 9, 8))
        check("중복은 추가 안 됨", added2, 0)
        check("누적 건수 유지", len(rows2), len(src_notices))

        made = archive.build_xlsx(rows2, xlsxp)
        if made is None:
            print("  (openpyxl 미설치 — 엑셀 검증 건너뜀)")
        else:
            from openpyxl import load_workbook

            ws = load_workbook(xlsxp).active
            check("시트 이름", ws.title, "공고누적")
            check("데이터 행 수", ws.max_row - 1, len(src_notices))
            check(
                "헤더",
                [c.value for c in ws[1]],
                ["발견일", "기관", "공고명", "게시일", "마감일", "매칭 키워드", "링크"],
            )
            check("'키' 컬럼은 제외", "키" in [c.value for c in ws[1]], False)
            check("고정틀", ws.freeze_panes, "A2")
            painted = sum(
                1
                for row in ws.iter_rows()
                for c in row
                if c.fill and c.fill.fgColor and c.fill.fgColor.rgb not in (None, "00000000")
            )
            check("색 채운 셀 없음", painted, 0)
            check("링크는 하이퍼링크", ws.cell(2, 7).value, "바로가기")

    print("\n[10-a2] 오늘치 엑셀 — 누적본과 양식이 같은지")
    import tempfile as _tf
    from pathlib import Path as _P

    from src import archive as _ar

    src_n = P(fixtures.BTP, inst(id="btp", name="부산테크노파크", base="https://www.btp.or.kr/"))
    for n in src_n:
        n.matched_keywords = ["테스트"]

    with _tf.TemporaryDirectory() as td:
        out = _P(td)
        daily = _ar.build_xlsx(
            _ar.rows_of(src_n, TODAY), _ar.daily_xlsx_path(out, TODAY), sheet_title="공고브리핑"
        )
        cum_rows, _ = _ar.add([], src_n, TODAY)
        cum = _ar.build_xlsx(cum_rows, out / "공고누적.xlsx")
        if daily is None or cum is None:
            print("  (openpyxl 미설치 — 건너뜀)")
        else:
            from openpyxl import load_workbook as _lw

            a, b = _lw(daily).active, _lw(cum).active
            check("파일명에 날짜", daily.name, "공고브리핑_2026-09-07.xlsx")
            check("시트명(오늘치)", a.title, "공고브리핑")
            check("시트명(누적)", b.title, "공고누적")
            check(
                "헤더 동일",
                [c.value for c in a[1]],
                [c.value for c in b[1]],
            )
            check(
                "열 너비 동일",
                [a.column_dimensions[k].width for k in "ABCDEFG"],
                [b.column_dimensions[k].width for k in "ABCDEFG"],
            )
            check("고정틀 동일", a.freeze_panes, b.freeze_panes)
            check("행 수 동일", a.max_row, b.max_row)
            painted = sum(
                1
                for ws in (a, b)
                for r in ws.iter_rows()
                for c in r
                if c.fill and c.fill.fgColor and c.fill.fgColor.rgb not in (None, "00000000")
            )
            check("색 채운 셀 없음", painted, 0)

    print("\n[10-b] 누적 엑셀 갱신 요일 (월·수·금)")
    from src.config import Config as _Cfg

    cfg_w = load_config()
    check("기본 요일", cfg_w.archive_weekdays, [0, 2, 4])
    for d, want in [
        (date(2026, 9, 7), True),    # 월
        (date(2026, 9, 8), False),   # 화
        (date(2026, 9, 9), True),    # 수
        (date(2026, 9, 10), False),  # 목
        (date(2026, 9, 11), True),   # 금
    ]:
        check(f"{d} 갱신 여부", d.weekday() in cfg_w.archive_weekdays, want)
    check("화요일의 다음 갱신일", cfg_w.next_archive_day(date(2026, 9, 8)), "수요일")
    check("금요일의 다음 갱신일", cfg_w.next_archive_day(date(2026, 9, 11)), "월요일")
    check("예전 단수 설정도 동작", _Cfg(archive_attach_weekday=0).archive_weekdays, [0])

    # ---------------------------------------------------------------- 구글시트
    print("\n[11] 구글시트 연동")
    from src import sheet as S

    INC_CSV = (
        "키워드,사용,메모\n"
        "시니어,Y,\n"
        "에이지테크,,비었으면 사용으로 봄\n"
        "쓰지않음,N,꺼둔 줄\n"
        ",,빈 줄\n"
    )
    EXC_CSV = (
        "키워드,사용,메모\n"
        "채용,Y,\n"
        " 입찰 ,Y,앞뒤 공백\n"
        "안쓰는제외,N,\n"
    )
    KW_CSV = (
        "구분,키워드,사용,메모\n"
        "포함,시니어,Y,\n"
        "포함,에이지테크,,비었으면 사용으로 봄\n"
        "포함,쓰지않음,N,꺼둔 줄\n"
        "제외,채용,Y,\n"
        "제외, 입찰 ,Y,앞뒤 공백\n"
        ",,,빈 줄\n"
    )
    RC_CSV = (
        "이메일,이름,사용,메모\n"
        "a@zeroweb.co.kr,홍길동,Y,\n"
        "b@zeroweb.co.kr,김철수,N,휴직\n"
        "잘못된주소,이영희,Y,@ 없음\n"
    )
    IN_CSV = (
        "기관명,주소,사용,메모\n"
        "부산정보산업진흥원,,N,잠시 끔\n"
        "KOTRA,,Y,다시 켬\n"
        "새로운진흥원,https://example.or.kr/notice,Y,신규\n"
        "이름만있고주소없음,,Y,추가되면 안 됨\n"
    )

    import csv as _csv
    import io as _io

    FAKE = {"kw": KW_CSV, "inc": INC_CSV, "exc": EXC_CSV, "rc": RC_CSV, "in": IN_CSV}

    def fake_get(url, timeout):
        """네트워크 대신 위 CSV 문자열을 돌려준다 (구글시트 게시본 형태)."""
        return list(_csv.DictReader(_io.StringIO(FAKE[url])))

    S._get = fake_get

    cfg2 = load_config()
    cfg2.sheet = {
        "enabled": True,
        "include_keywords_url": "inc",
        "exclude_keywords_url": "exc",
        "recipients_url": "rc",
        "institutions_url": "in",
    }
    insts = load_institutions_for_test()
    insts, log = S.apply_all(cfg2, insts)

    check("포함 키워드", cfg2.include_keywords, ["시니어", "에이지테크"])
    check("제외 키워드", cfg2.exclude_keywords, ["채용", "입찰"])
    check("수신자 (N·잘못된 주소 제외)", cfg2.recipients, ["a@zeroweb.co.kr"])
    by = {i.name: i for i in insts}
    check("시트에서 끈 기관", by["부산정보산업진흥원"].enabled, False)
    check("시트에서 켠 기관", by["KOTRA"].enabled, True)
    check("새 기관 추가됨", "새로운진흥원" in by, True)
    check("주소 없는 새 이름은 무시", "이름만있고주소없음" in by, False)
    check("로그 4줄", len(log), 4)

    print("\n[11-a2] 예전 방식(한 탭에 '구분' 칸)도 계속 동작")
    cfg2b = load_config()
    cfg2b.sheet = {"enabled": True, "keywords_url": "kw"}
    S.apply_all(cfg2b, load_institutions_for_test())
    check("포함 키워드", cfg2b.include_keywords, ["시니어", "에이지테크"])
    check("제외 키워드", cfg2b.exclude_keywords, ["채용", "입찰"])

    print("\n[11-b] 시트를 못 읽어도 저장소 설정으로 계속 진행")

    def boom(url, timeout):
        raise S.SheetError("HTTP 404")

    S._get = boom
    cfg3 = load_config()
    before_inc = list(cfg3.include_keywords)
    cfg3.sheet = {"enabled": True, "keywords_url": "kw", "recipients_url": "rc"}
    insts3, log3 = S.apply_all(cfg3, load_institutions_for_test())
    check("키워드 그대로", cfg3.include_keywords, before_inc)
    check("실패가 로그에 남음", all("실패" in x for x in log3), True)
    check("기관 목록 유지", len(insts3) > 30, True)

    print("\n[11-b2] 편집 주소를 CSV 주소로 자동 변환")
    SID = "1Mf3bRkTTbl-LhRSBsqcP1zOTZe4pN4A9Fn2EsP7aZko"
    check(
        "편집 주소 → export CSV",
        S.to_csv_url(f"https://docs.google.com/spreadsheets/d/{SID}/edit?gid=441632177#gid=441632177"),
        f"https://docs.google.com/spreadsheets/d/{SID}/export?format=csv&gid=441632177",
    )
    check(
        "gid 없으면 첫 탭(0)",
        S.to_csv_url(f"https://docs.google.com/spreadsheets/d/{SID}/edit"),
        f"https://docs.google.com/spreadsheets/d/{SID}/export?format=csv&gid=0",
    )
    pub = "https://docs.google.com/spreadsheets/d/e/2PACX-1vABC/pub?gid=0&single=true&output=csv"
    check("웹에 게시 주소는 그대로", S.to_csv_url(pub), pub)
    check("빈 값은 빈 값", S.to_csv_url(""), "")

    print("\n[11-c] 시트를 끄면 아무것도 안 함")
    cfg4 = load_config()
    cfg4.sheet = {"enabled": False, "keywords_url": "kw"}
    _, log4 = S.apply_all(cfg4, load_institutions_for_test())
    check("로그 없음", log4, [])

    print("\n[13] 외부 배너·바로가기를 공고로 착각하지 않는지 (한국사회적기업진흥원)")
    # 메인 페이지를 목록으로 쓰면 '판로지원 플랫폼 → sepp.or.kr' 같은
    # 외부 배너가 공고로 딸려 들어왔다. 실제로 메일에 2건 나갔다.
    banner_html = """<html><body><table><tbody>
      <tr><td>1</td><td><a href="https://www.sepp.or.kr/store365">판로지원 플랫폼</a></td>
          <td>2026-09-08</td></tr>
      <tr><td>2</td><td><a href="https://www.coop.go.kr/home/index.do">협동조합 홍보포털</a></td>
          <td>2026-09-08</td></tr>
      <tr><td>3</td><td><a href="/homepage/bbs/boardView.do?bIdx=1">
          2026년 사회적기업 인증 지원사업 공고</a></td><td>2026-09-08</td></tr>
      <tr><td>4</td><td><a href="/homepage/bbs/boardView.do?bIdx=2">
          「2026년 제2차 사회서비스 정책포럼」 개최 안내 2026-09-08 23</a></td><td>2026-09-08</td></tr>
      <tr><td>5</td><td><a href="/homepage/bbs/boardView.do?bIdx=3">
          2026 서울창업허브 공덕 9월 허브아워 새로운게시글</a></td><td>2026-09-03</td></tr>
    </tbody></table></body></html>"""
    r = P(
        banner_html,
        inst(
            id="se",
            name="한국사회적기업진흥원",
            url="https://socialenterprise.or.kr/homepage/main.do",
            base="https://socialenterprise.or.kr",
        ),
    )
    titles13 = [n.title for n in r]
    check("외부 배너 2건 제외", len(r), 3)
    check("'판로지원 플랫폼' 없음", "판로지원 플랫폼" in titles13, False)
    check("'협동조합 홍보포털' 없음", "협동조합 홍보포털" in titles13, False)
    check("본문 공고는 남음", "2026년 사회적기업 인증 지원사업 공고" in titles13, True)

    print("\n[13-a2] 두 span 내용이 다를 때도 제목 중복 제거 (2026-09-09 재발분)")
    from src.parse import _collapse_doubled as _CD

    for src_t, want_t in [
        (
            "2026년 수리조선 기업인증 지원사업 추가모집 공고 재공고 2026년 수리조선 기업인증 지원사업 추가모집 공고",
            "2026년 수리조선 기업인증 지원사업 추가모집 공고 재공고",
        ),
        (
            "기업성장기반 글로벌 하이메디 허브 특구 상생협력사업 공고(4차) 재공고 기업성장기반 글로벌 하이메디 허브 특구 상생협...",
            "기업성장기반 글로벌 하이메디 허브 특구 상생협력사업 공고(4차) 재공고",
        ),
        ("신중년 디지털 전환 지원사업 공고 신중년 디지털 전환 지원사업 공고", "신중년 디지털 전환 지원사업 공고"),
        # 정상 제목은 건드리지 않는다
        ("2026년 AX 에이지테크 시장확산 지원사업 모집 공고", "2026년 AX 에이지테크 시장확산 지원사업 모집 공고"),
        (
            "[공고] 「HLTH USA 2026」 디지털헬스케어 한국관 참가기관 모집 공고(~9.30.(수)) 15시까지",
            "[공고] 「HLTH USA 2026」 디지털헬스케어 한국관 참가기관 모집 공고(~9.30.(수)) 15시까지",
        ),
        (
            "2026년 사회서비스 종사자 심리상담 지원사업 3기 참여자 모집(~9/30)",
            "2026년 사회서비스 종사자 심리상담 지원사업 3기 참여자 모집(~9/30)",
        ),
    ]:
        check(f"'{src_t[:30]}…'", _CD(src_t), want_t)

    print("\n[13-a3] 텍스트 메일에서 '마감 마감일 확인필요' 중복 표현 제거")
    from src.render import render_text as _rt
    from src.parse import Notice as _N

    _cfg = load_config()
    _g = {
        "x": [
            _N("x", "테스트기관", "마감일 없는 공고", "https://x.kr/1", date(2026, 9, 8), None),
            _N("x", "테스트기관", "상시모집 공고", "https://x.kr/2", date(2026, 9, 8), None, False, [], "상시"),
        ]
    }
    txt = _rt(_g, {"x": "테스트기관"}, TODAY, _cfg)
    check("'마감 마감일' 중복 없음", "마감 마감일" in txt, False)
    check("'마감 확인필요' 로 표기", "마감 확인필요" in txt, True)
    check("상시모집 표기", "마감 상시모집" in txt, True)

    print("\n[13-b] 제목 뒤 '등록일+조회수' 꼬리와 '새로운게시글' 배지 제거")
    check("등록일·조회수 꼬리 제거", titles13[1], "「2026년 제2차 사회서비스 정책포럼」 개최 안내")
    check("'새로운게시글' 제거", titles13[2], "2026 서울창업허브 공덕 9월 허브아워")

    print("\n[13-c] 서브도메인이 달라도 같은 기관이면 유지 (전북테크노파크)")
    from src.parse import _same_site

    for u, b, want in [
        ("https://www.jbtp.or.kr/board/view.jbtp?x=1", "https://jbcis.jbtp.or.kr", True),
        ("https://www.sepp.or.kr/store365", "https://socialenterprise.or.kr", False),
        ("https://www.coop.go.kr/home", "https://socialenterprise.or.kr", False),
        ("https://bipa.kr/board/x", "https://bipa.kr", True),
        ("/relative/path", "https://bipa.kr", True),
    ]:
        check(f"{u[:38]}", _same_site(u, b), want)

    # ---------------------------------------------------------------- 마감일 판정
    print("\n[12] 마감일 읽기 — 실제 공고에서 쓰이는 표기들")
    from src import deadline as DL

    for text, want_dl, want_note in [
        ("2026년 GHKOL 컨설팅 사업 공고 (~9.16.(수) 까지)", date(2026, 9, 16), ""),
        ("에이지테크 시장확산 지원사업 모집 (~2026.9.30.)", date(2026, 9, 30), ""),
        ("접수기간 : 2026.09.07 ~ 2026.09.30", date(2026, 9, 30), ""),
        ("신청기간 2026-09-07 ~ 2026-10-06", date(2026, 10, 6), ""),
        ("접수마감: 2026. 9. 30.(월) 18:00", date(2026, 9, 30), ""),
        ("마감 2026-12-31 18:00", date(2026, 12, 31), ""),
        ("2026년 9월 30일까지 신청", date(2026, 9, 30), ""),
        ("입주기업 상시모집 공고(3차)", None, "상시"),
        ("예산 소진 시까지 선착순 접수", None, "상시"),
        ("2026년 1차 고령친화우수제품 지정 공고", None, ""),
    ]:
        check(f"'{text[:26]}…'", DL.from_text(text, TODAY), (want_dl, want_note))

    print("\n[12-a2] 연도 없는 날짜의 해 추론 — 기준일에 가장 가까운 해를 고른다")
    # 실제로 새어나간 공고: 9월에 받은 메일에 '(~4.24.(금))' 공고가 들어 있었다.
    # 예전 규칙('과거면 내년')은 이걸 2027-04-24 로 보고 진행 중이라 판단했다.
    slipped = (
        "AI응용제품 신속 상용화(복지분야, 에이지테크 기반 고령친화사업 지원) "
        "컨소시엄 모집 공고(~4.24.(금), 18:00)"
    )
    check("실제 새어나간 공고의 마감일", DL.from_text(slipped, TODAY)[0], date(2026, 4, 24))
    check("→ 오늘보다 과거", DL.from_text(slipped, TODAY)[0] < TODAY, True)

    # 기준일이 '오늘'일 때 (게시일 모르는 게시판) — 오늘과 가장 가까운 해
    for ref, txt, want in [
        (date(2026, 9, 7), "(~4.24.(금), 18:00)", date(2026, 4, 24)),
        (date(2026, 9, 7), "(~9.16.(수) 까지)", date(2026, 9, 16)),
        (date(2026, 12, 20), "(~1.15. 까지)", date(2027, 1, 15)),
    ]:
        check(f"오늘 {ref} · '{txt}'", DL.from_text(txt, ref)[0], want)

    # 기준일이 '게시일'일 때 — 마감일은 게시일보다 앞설 수 없다
    for ref, txt, want in [
        (date(2026, 1, 5), "(~12.20. 까지)", date(2026, 12, 20)),
        (date(2026, 4, 2), "(~4.24.(금))", date(2026, 4, 24)),
        (date(2026, 12, 28), "(~1.15. 까지)", date(2027, 1, 15)),
        (date(2026, 9, 5), "(~9.30.)", date(2026, 9, 30)),
    ]:
        check(
            f"게시 {ref} · '{txt}'",
            DL.from_text(txt, ref, anchor_is_posted=True)[0],
            want,
        )

    print("\n[12-a3] 게시일을 알면 그걸 기준으로 해를 정한다")
    html_old = """<html><body><table><tbody>
      <tr><td>1</td><td><a href="/v?1">에이지테크 컨소시엄 모집 공고(~4.24.(금), 18:00)</a></td>
          <td>2026-04-02</td></tr>
      <tr><td>2</td><td><a href="/v?2">시니어 돌봄 실증 참여기업 모집(~9.30.)</a></td>
          <td>2026-09-05</td></tr>
      <tr><td>3</td><td><a href="/v?3">고령친화 우수제품 지정 공고</a></td>
          <td>2026-09-01</td></tr>
    </tbody></table></body></html>"""
    r = P(html_old, inst(id="t3", base="https://example.com"))
    check("게시일 기준 해 추론", r[0].deadline, date(2026, 4, 24))
    check("두 번째 공고", r[1].deadline, date(2026, 9, 30))

    print("\n[12-a4] 게시일이 없는 게시판에서도 마감된 공고는 걸러진다")
    html_nodate = """<html><body><table class="tstyle_list"><tbody>
      <tr><td>410</td><td><a href="/board/view?linkId=48942096">
        AI응용제품 신속 상용화(복지분야) 컨소시엄 모집 공고(~4.24.(금), 18:00)</a></td>
        <td>320</td><td>첨부</td></tr>
      <tr><td>416</td><td><a href="/board/view?linkId=48948098">
        2026년 1차 고령친화우수제품 지정 공고</a></td><td>211</td><td>첨부</td></tr>
      <tr><td>417</td><td><a href="/board/view?linkId=48948200">
        고령친화 실증 참여기업 모집(~12.31.)</a></td><td>187</td><td>첨부</td></tr>
    </tbody></table></body></html>"""
    r = P(html_nodate, inst(id="khidi_e", base="https://www.khidi.or.kr"))
    check("게시일 없음", r[0].posted, None)
    check("그래도 마감일은 잡힘", r[0].deadline, date(2026, 4, 24))

    cfg_e = load_config()
    cfg_e.exclude_expired = True
    cfg_e.unknown_deadline = "include_flagged"
    kept_e, dropped_e = drop_expired_fn(r, cfg_e, TODAY)
    titles_e = [n.title for n in kept_e]
    check("4월 마감 공고 제외됨", any("4.24" in t for t in titles_e), False)
    check("12월 마감 공고 유지", any("12.31" in t for t in titles_e), True)
    check("마감일 미상 공고 유지", any("우수제품 지정" in t for t in titles_e), True)

    print("\n[12-b] 마감 상태 판정 — '마감임박'을 마감으로 착각하지 않는지")
    for text, want in [
        ("접수마감", True),
        ("모집 종료", True),
        ("접수가 마감되었습니다", True),
        ("마감임박 D-3", False),
        ("마감일 2026-09-30", False),
        ("마감기한 : D-4", False),
        ("접수중", False),
    ]:
        check(f"'{text}'", DL.is_closed(text), want)

    print("\n[12-c] 상세 페이지에서 마감일 읽기")
    detail = """<html><body><nav>메뉴</nav><div class="view">
      <h2>2026년 스마트서비스 지원사업 참여기업 모집 공고</h2>
      <table><tr><th>접수기간</th><td>2026-09-01 ~ 2026-09-05</td></tr>
      <tr><th>담당부서</th><td>기업지원단</td></tr></table>
      <p>자세한 내용은 첨부파일을 참고하시기 바랍니다.</p>
    </div></body></html>"""
    dl, note, closed = DL.from_detail_html(detail, TODAY)
    check("상세에서 마감일", dl, date(2026, 9, 5))
    check("상시 아님", note, "")
    check("마감문구 없음", closed, False)

    print("\n[12-d] 마감 확인된 공고가 실제로 걸러지는지")
    from src.verify import drop_expired
    from src.parse import Notice

    cfg5 = load_config()
    cfg5.exclude_expired = True
    cfg5.unknown_deadline = "include_flagged"
    items = [
        Notice("a", "A", "아직 접수중인 공고", "u1", date(2026, 9, 5), date(2026, 9, 30)),
        Notice("a", "A", "어제 마감된 공고", "u2", date(2026, 9, 1), date(2026, 9, 6)),
        Notice("a", "A", "오늘 마감인 공고", "u3", date(2026, 9, 1), date(2026, 9, 7)),
        Notice("a", "A", "마감 문구가 있는 공고", "u4", date(2026, 9, 1), None, True),
        Notice("a", "A", "마감일 모르는 공고", "u5", date(2026, 9, 5), None),
        Notice("a", "A", "상시모집 공고", "u6", date(2026, 9, 5), None, False, [], "상시"),
    ]
    kept, dropped = drop_expired(items, cfg5, TODAY)
    titles = [n.title for n in kept]
    check("어제 마감 → 제외", "어제 마감된 공고" in titles, False)
    check("마감 문구 → 제외", "마감 문구가 있는 공고" in titles, False)
    check("오늘 마감 → 포함", "오늘 마감인 공고" in titles, True)
    check("접수중 → 포함", "아직 접수중인 공고" in titles, True)
    check("상시모집 → 포함", "상시모집 공고" in titles, True)
    check("마감일 미상 → 포함(확인필요)", "마감일 모르는 공고" in titles, True)
    check("제외된 수", dropped, 2)

    print("\n[12-e] unknown_deadline=exclude 로 바꾸면 미상도 제외 (상시는 유지)")
    cfg5.unknown_deadline = "exclude"
    kept2, dropped2 = drop_expired(items, cfg5, TODAY)
    t2 = [n.title for n in kept2]
    check("마감일 미상 → 제외", "마감일 모르는 공고" in t2, False)
    check("상시모집은 유지", "상시모집 공고" in t2, True)

    print("\n[12-f] 제목에 마감일이 있으면 파서가 바로 잡는지")
    html = """<html><body><table><tbody>
      <tr><td>1</td><td><a href="/v?1">2026년 컨설팅 지원사업 공고 (~9.16.(수) 까지)</a></td>
          <td>2026-09-03</td></tr>
      <tr><td>2</td><td><a href="/v?2">고령친화 실증 참여기업 상시모집</a></td>
          <td>2026-09-02</td></tr>
      <tr><td>3</td><td><a href="/v?3">시니어 돌봄 지원사업 모집 공고</a></td>
          <td>2026-09-01</td></tr>
    </tbody></table></body></html>"""
    r = P(html, inst(id="t2", base="https://example.com"))
    check("제목에서 마감일", r[0].deadline, date(2026, 9, 16))
    check("출처 표시", r[0].deadline_source, "제목")
    check("상시모집 인식", r[1].note, "상시")
    check("마감일 없는 건 그대로", r[2].deadline, None)

    print()
    if FAILS:
        print(f"실패 {len(FAILS)}건")
        for f in FAILS:
            print("  - " + f)
        return 1
    print("모든 검증 통과")
    return 0


if __name__ == "__main__":
    sys.exit(main())
