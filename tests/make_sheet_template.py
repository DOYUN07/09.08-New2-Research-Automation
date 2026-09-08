"""구글시트에 올릴 설정시트 템플릿을 만든다.

실행: python -m tests.make_sheet_template   →  out/공고브리핑_설정시트.xlsx

만들어진 파일을 구글 드라이브에 올리고
  파일 → Google 스프레드시트로 저장
하면 탭 3개가 그대로 들어간 시트가 됩니다. 색은 쓰지 않습니다.
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from src.config import ROOT, load_config, load_institutions

FONT = "맑은 고딕"


def _style(ws, widths: list[int], note: str) -> None:
    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in range(1, len(widths) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, size=10, bold=True)
        cell.border = box
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.border = box
            cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    # 안내문은 표 오른쪽에 둔다 (CSV로 읽을 때 방해되지 않도록 열을 띄움)
    col = len(widths) + 2
    ws.cell(row=1, column=col, value="사용법").font = Font(name=FONT, size=10, bold=True)
    for i, line in enumerate(note.strip().split("\n"), start=2):
        ws.cell(row=i, column=col, value=line.strip()).font = Font(name=FONT, size=9)
    ws.column_dimensions[get_column_letter(col)].width = 60


def build(path, include_recipients: bool = False):
    """설정시트를 만든다.

    include_recipients=False 가 기본이다. '웹에 게시'한 시트는 주소를 아는 사람이면
    누구나 볼 수 있어서, 수신자 이메일은 시트에 두지 않고 config.yaml 로 관리하는 편이
    안전하기 때문이다.
    """
    cfg = load_config()
    wb = Workbook()

    # ---------------- 포함키워드 ----------------
    ws = wb.active
    ws.title = "포함키워드"
    ws.append(["키워드", "사용", "메모"])
    for w in cfg.include_keywords:
        ws.append([w, "Y", ""])
    ws.append(["", "Y", "← 여기에 한 줄씩 추가하세요"])
    _style(
        ws,
        [26, 8, 34],
        """
        공고 제목에 이 말이 하나라도 있으면 메일에 넣습니다.

        사용: N 이라고 적으면 그 줄은 무시합니다. 지우지 말고 N 으로 꺼두세요.
        줄 추가는 맨 아래에 이어서 쓰면 됩니다.

        공고가 너무 안 오면 → 여기에 키워드를 더 넣으세요.
        너무 많이 오면 → 넓은 단어(예: '안전', '홍보')를 N 으로 꺼보세요.

        ※ 제외키워드가 포함키워드보다 우선합니다.
          양쪽에 다 걸리는 공고는 버려집니다.
        """,
    )

    # ---------------- 제외키워드 ----------------
    ws = wb.create_sheet("제외키워드")
    ws.append(["키워드", "사용", "메모"])
    for w in cfg.exclude_keywords:
        ws.append([w, "Y", ""])
    ws.append(["", "Y", "← 여기에 한 줄씩 추가하세요"])
    _style(
        ws,
        [26, 8, 34],
        """
        공고 제목에 이 말이 있으면 무조건 버립니다.
        포함키워드에 걸려도 여기 걸리면 버려집니다.

        사용: N 이라고 적으면 그 줄은 무시합니다.
        줄 추가는 맨 아래에 이어서 쓰면 됩니다.

        쓸데없는 공고가 자꾸 오면 → 그 공고 제목에서 공통된 말을 여기에 넣으세요.
        (예: 채용, 입찰, 용역, 결과 발표)

        ※ 너무 흔한 말은 넣지 마세요.
          '지원' 같은 단어를 넣으면 필요한 공고까지 다 사라집니다.
        """,
    )

    # ---------------- 수신자 (기본으로는 넣지 않음) ----------------
    if include_recipients:
        ws = wb.create_sheet("수신자")
        ws.append(["이메일", "이름", "사용", "메모"])
        for m in cfg.recipients:
            ws.append([m, "", "Y", ""])
        ws.append(["", "", "Y", "← 여기에 한 줄씩 추가하세요"])
        _style(
            ws,
            [34, 14, 8, 30],
            """
            메일 받을 사람을 한 줄에 하나씩 적습니다.
            사용: N 이면 그 사람에게는 안 보냅니다.

            ⚠️ 시트를 '웹에 게시'하면 주소를 아는 사람은 누구나 볼 수 있습니다.
            """,
        )

    # ---------------- 기관 ----------------
    ws = wb.create_sheet("기관")
    ws.append(["기관명", "주소", "사용", "메모"])
    for inst in load_institutions():
        ws.append(
            [
                inst.name,
                "" if inst.status in ("ok", "check") else inst.url,
                "Y" if inst.enabled else "N",
                (inst.notes or "").strip().split("\n")[0][:60],
            ]
        )
    _style(
        ws,
        [30, 44, 8, 62],
        """
        기존 기관은 사용 칸만 Y / N 으로 바꾸면 켜고 꺼집니다.
        주소 칸은 비워두세요 — 세부 설정은 저장소가 계속 관리합니다.

        새 기관을 넣고 싶으면 맨 아래에 기관명 + 주소(공고 목록 페이지)를 쓰고
        사용에 Y 를 적으면 다음 실행부터 수집합니다.

        기관명을 고치면 기존 기관과 연결이 끊어지니 그대로 두세요.
        """,
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    import sys

    with_rcpt = "--with-recipients" in sys.argv
    p = build(ROOT / "out" / "공고브리핑_설정시트.xlsx", include_recipients=with_rcpt)
    from openpyxl import load_workbook

    tabs = " · ".join(load_workbook(p).sheetnames)
    print(f"만들었습니다: {p}  (탭: {tabs})")
    if not with_rcpt:
        print("수신자 탭까지 넣으려면: python -m tests.make_sheet_template --with-recipients")
