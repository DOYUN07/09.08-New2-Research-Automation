"""공고 누적 기록.

두 개의 파일을 씁니다.

  state/archive.csv   실제 저장소(원본). 매일 갱신되고 GitHub에 커밋됩니다.
                      텍스트라 변경 내역이 그대로 보이고 용량도 거의 안 늘어납니다.

  공고누적.xlsx        보기용. CSV에서 매번 새로 만듭니다.
                      저장소 최상단에 놓이고, 주 1회 메일에도 첨부됩니다.

엑셀을 원본으로 두지 않는 이유: 엑셀은 이진 파일이라 매일 커밋하면 저장소가
빠르게 무거워지고, 변경 내역도 볼 수 없습니다.
"""
from __future__ import annotations

import csv
from datetime import date
from pathlib import Path

from .config import ROOT
from .parse import Notice

ARCHIVE_CSV = ROOT / "state" / "archive.csv"
ARCHIVE_XLSX = ROOT / "공고누적.xlsx"

HEADERS = ["발견일", "기관", "공고명", "게시일", "마감일", "매칭 키워드", "링크", "키"]


def _fmt(d: date | None) -> str:
    return d.isoformat() if d else ""


def load_rows(path: Path | None = None) -> list[dict]:
    p = path or ARCHIVE_CSV
    if not p.exists():
        return []
    try:
        with p.open(encoding="utf-8-sig", newline="") as fh:
            return [r for r in csv.DictReader(fh) if r.get("키")]
    except OSError:
        return []


def add(rows: list[dict], notices: list[Notice], found: date) -> tuple[list[dict], int]:
    """오늘 발송한 공고를 누적 목록에 더한다. (전체 행, 새로 추가된 수)"""
    known = {r["키"] for r in rows}
    added = 0
    for n in notices:
        if n.key in known:
            continue
        rows.append(
            {
                "발견일": found.isoformat(),
                "기관": n.institution_name,
                "공고명": n.title,
                "게시일": _fmt(n.posted),
                "마감일": _fmt(n.deadline) or "확인필요",
                "매칭 키워드": ", ".join(n.matched_keywords),
                "링크": n.url,
                "키": n.key,
            }
        )
        known.add(n.key)
        added += 1
    return rows, added


def save_csv(rows: list[dict], path: Path | None = None) -> None:
    p = path or ARCHIVE_CSV
    p.parent.mkdir(parents=True, exist_ok=True)
    # utf-8-sig: 이 파일을 엑셀에서 바로 열어도 한글이 깨지지 않게
    with p.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=HEADERS)
        w.writeheader()
        w.writerows(sorted(rows, key=lambda r: (r["발견일"], r["기관"]), reverse=True))


def build_xlsx(rows: list[dict], path: Path) -> Path | None:
    """누적 CSV를 보기 좋은 엑셀로 만든다. 색은 쓰지 않는다(흑백 기본 서식)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None

    ordered = sorted(rows, key=lambda r: (r["발견일"], r["게시일"]), reverse=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "공고누적"

    base = Font(name="맑은 고딕", size=10)
    head = Font(name="맑은 고딕", size=10, bold=True)
    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    visible = HEADERS[:-1]  # '키'는 내부용이라 엑셀에는 넣지 않는다
    ws.append(visible)
    for c in range(1, len(visible) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = head
        cell.border = box
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in ordered:
        ws.append([r.get(h, "") for h in visible])

    link_col = visible.index("링크") + 1
    for row in range(2, ws.max_row + 1):
        for col in range(1, len(visible) + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = base
            cell.border = box
            cell.alignment = Alignment(vertical="top", wrap_text=(col == 3))
        url = ws.cell(row=row, column=link_col).value
        if url:
            c = ws.cell(row=row, column=link_col)
            c.hyperlink = url
            c.value = "바로가기"
            c.alignment = Alignment(horizontal="center", vertical="top")

    # 발견일 / 기관 / 공고명 / 게시일 / 마감일 / 매칭 키워드 / 링크
    widths = [12, 24, 62, 12, 12, 28, 10]
    for i, w in enumerate(widths[: len(visible)], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(visible))}{max(ws.max_row, 1)}"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
